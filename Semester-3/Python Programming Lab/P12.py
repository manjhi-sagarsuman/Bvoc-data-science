# Practical 12: Student Details System Using Excel
from openpyxl import Workbook, load_workbook
import os

FILE = "student_details.xlsx"
SHEET = "Students"
# Create Excel workbook if it does not exist
if not os.path.exists(FILE):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    sheet.append(["Roll No", "Name", "Course", "Age"])
    workbook.save(FILE)

def create_student():
    roll = input("Enter Roll Number: ").strip()
    name = input("Enter Student Name: ").strip()
    course = input("Enter Course: ").strip()
    age = input("Enter Age: ").strip()

    if not all([roll, name, course, age]):
        print("All fields are required.")
        return

    if not age.isdigit() or int(age) <= 0:
        print("Invalid age.")
        return

    wb = load_workbook(FILE)
    ws = wb[SHEET]

    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[0]) == roll:
            print("Roll number already exists.")
            wb.close()
            return

    ws.append([roll, name, course, int(age)])
    wb.save(FILE)
    wb.close()
    print("Student added successfully.")

def read_students():
    wb = load_workbook(FILE, read_only=True, data_only=True)
    ws = wb[SHEET]

    print("\n--- Student Details ---")
    for row in ws.iter_rows(values_only=True):
        print(" | ".join(str(value) for value in row))

    wb.close()

def update_student():
    roll = input("Enter Roll Number to update: ").strip()

    wb = load_workbook(FILE)
    ws = wb[SHEET]

    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == roll:
            name = input("Enter new Name: ").strip()
            course = input("Enter new Course: ").strip()
            age = input("Enter new Age: ").strip()

            if not name or not course or not age.isdigit() or int(age) <= 0:
                print("Invalid details.")
                wb.close()
                return

            row[1].value = name
            row[2].value = course
            row[3].value = int(age)

            wb.save(FILE)
            wb.close()
            print("Student updated successfully.")
            return

    wb.close()
    print("Student not found.")

def delete_student():
    roll = input("Enter Roll Number to delete: ").strip()

    wb = load_workbook(FILE)
    ws = wb[SHEET]

    for row_num in range(2, ws.max_row + 1):
        if str(ws.cell(row=row_num, column=1).value) == roll:
            ws.delete_rows(row_num)
            wb.save(FILE)
            wb.close()
            print("Student deleted successfully.")
            return

    wb.close()
    print("Student not found.")
# Main menu
while True:
    print("\n===== Student Details System =====")
    print("1. Create Student")
    print("2. Read Student Details")
    print("3. Update Student")
    print("4. Delete Student")
    print("5. Exit")

    choice = input("Enter your choice: ").strip()

    if choice == "1":
        create_student()
    elif choice == "2":
        read_students()
    elif choice == "3":
        update_student()
    elif choice == "4":
        delete_student()
    elif choice == "5":
        print("Thank you for using the Student Details System.")
        break
    else:
        print("Invalid choice. Please select 1-5.")
